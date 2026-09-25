"""Экспорт табличных отчётов в CSV (для Excel) и DOCX (для печати/согласования)."""
import csv
import io
from datetime import date
from typing import Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def to_csv(headers: Sequence[str], rows: Sequence[Sequence]) -> bytes:
    # «;» и BOM — чтобы русский Excel открыл файл двойным щелчком без мастера импорта
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow(headers)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return buf.getvalue().encode("utf-8-sig")


def to_docx(title: str, headers: Sequence[str], rows: Sequence[Sequence], subtitle: str = "",
            landscape: bool = True, col_widths_cm: Sequence[float] | None = None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title)
    run.bold = True
    run.font.size = Pt(14)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run((subtitle + " · " if subtitle else "") + f"по состоянию на {date.today():%d.%m.%Y}").italic = True

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(text))
        r.bold = True
        r.font.size = Pt(11)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11)
    if col_widths_cm:
        for row in table.rows:
            for i, w in enumerate(col_widths_cm):
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    doc.add_paragraph(f"Всего записей: {len(rows)}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
